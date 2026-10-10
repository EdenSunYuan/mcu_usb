#!/usr/bin/env python3
"""
gen_report_docx.py — 把仓库里的 Markdown 报告转成 Word 版 (.docx)

取代原先一份报告一个的 `gen_<模块>_design_docx.py`（内容写死在 Python 里，报告一改脚本就过时）。
本脚本直接解析 Markdown 源，排版沿用同一套约定：宋体正文 / Consolas 代码、
蓝色分级标题、Table Grid 表头带底色、代码块带灰底、图居中内嵌 + 图注。

支持的 Markdown 子集（够本仓库的设计/验证报告用）：
    # 标题 / ## N. x / ### N.N x     四级标题
    > 行                             灰小字（报告头部的模块/日期/结论块）
    ``` 代码块 ```                    灰底等宽
    | 表格 | 且次行含 ---              Table Grid + 表头底色
    - 列表 / 1. 列表                  圆点缩进 / 原样编号
    **粗体** / `行内代码` / [文字](链接)  行内样式，链接转 Word 超链接
    📎 [图：标题（x.png）](../fig/x.png)  内嵌 PNG（相对 Markdown 所在目录解析）+ 图注
    ![alt](x.png)                    同上
    整行 *斜体*                       斜体灰字（报告末尾的注）

用法:
    python3 tools/gen_report_docx.py doc/design_rpt/bus_state_design.md
    python3 tools/gen_report_docx.py doc/design_rpt/bus_state_design.md -o doc/总线状态检测设计报告.docx
    python3 tools/gen_report_docx.py <md> --dry-run     # 只打印解析出的元素，不写文件
"""

import argparse
import os
import re

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BODY_FONT = "宋体"
MONO_FONT = "Consolas"
BODY_SIZE = 10.5
CODE_SIZE = 9.5
# 表格字号：列数一多就降一号，让整表仍放得下一页（xiao_yuan 2026-09-28）
NCOL_BODY = 5          # 列数 ≤5 用正文字号
NCOL_SMALL = 7         # 列数 ≤7 降到 9pt，更多降到 8pt
CODE_BG = "F2F2F2"
HDR_BG = "DEEAF6"
HEAD_COLOR = {1: "1F3864", 2: "2E5395", 3: "404040", 4: "404040"}
HEAD_SIZE = {1: 16, 2: 13, 3: 11.5, 4: 11}
META_COLOR = "595959"

DOCX_LINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"

# ── 行内样式 ────────────────────────────────────────────────────────────────


def set_run(run, size=BODY_SIZE, bold=False, mono=False, italic=False, color=None):
    run.font.name = MONO_FONT if mono else BODY_FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    if mono:
        rFonts.set(qn("w:ascii"), MONO_FONT)
        rFonts.set(qn("w:hAnsi"), MONO_FONT)
    rFonts.set(qn("w:eastAsia"), BODY_FONT)
    return run


def add_hyperlink(p, url, text, size=BODY_SIZE):
    part = p.part
    r_id = part.relate_to(url, DOCX_LINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    link.append(run)
    p._p.append(link)
    r = p.add_run(text)
    set_run(r, size=size, color="0563C1")
    run.addprevious(r._element)          # 样式化后挪进 hyperlink 节点内
    link.append(r._element)
    r.font.underline = True
    return r


def add_rich(p, text, size=BODY_SIZE, bold=False, color=None):
    """**粗体** / `行内代码` / [文字](链接) → 依次落到同一个段落。"""
    for tok in re.split(r"(\*\*.*?\*\*|`[^`]*`|\[[^\]]*\]\([^)]*\))", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            set_run(p.add_run(tok[2:-2]), size=size, bold=True, color=color)
        elif tok.startswith("`") and tok.endswith("`"):
            set_run(p.add_run(tok[1:-1]), size=size, mono=True, color=color)
        elif re.fullmatch(r"\[[^\]]*\]\([^)]*\)", tok):
            label, url = re.match(r"\[([^\]]*)\]\(([^)]*)\)", tok).groups()
            add_hyperlink(p, url, label, size=size)
        else:
            set_run(p.add_run(tok), size=size, bold=bold, color=color)


def add_para(doc, text="", size=BODY_SIZE, space_after=4, indent=None,
             italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    if italic:
        set_run(p.add_run(text), size=size, italic=True, color=color)
    else:
        add_rich(p, text, size=size, color=color)
    return p


def add_heading(doc, text, level):
    """用 Word 内置 Heading 样式（拿到大纲级别，导航窗格/目录可用），再覆盖字体与颜色。"""
    p = doc.add_paragraph()
    try:
        p.style = doc.styles[f"Heading {min(level, 4)}"]
    except KeyError:
        pass
    p.paragraph_format.space_before = Pt(14 if level == 1 else (10 if level == 2 else 8))
    p.paragraph_format.space_after = Pt(6)
    add_rich(p, text, size=HEAD_SIZE.get(level, 11), color=HEAD_COLOR.get(level, "404040"))
    for r in p.runs:
        r.bold = True
        if not r.font.color.rgb:
            r.font.color.rgb = RGBColor.from_string(HEAD_COLOR.get(level, "404040"))
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Pt(12)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), CODE_BG)
    p._p.get_or_add_pPr().append(shd)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        set_run(p.add_run(line), size=CODE_SIZE, mono=True)
    return p


def table_size(ncol):
    """表格字号按列数降档：列多了字就小一号，整表仍放得下一页。"""
    if ncol <= NCOL_BODY:
        return BODY_SIZE
    if ncol <= NCOL_SMALL:
        return 9
    return 8


def add_table(doc, header, rows, md_dir=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = True
    size = table_size(len(header))
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_rich(cell.paragraphs[0], h, size=size)
        for r in cell.paragraphs[0].runs:
            r.bold = True
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), HDR_BG)
        cell._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(row):
            if i >= len(header):
                break
            cells[i].text = ""
            full = None
            m = FIG_RE.match(c.strip())          # 单元格整格是一张图（如 TX 报告的波形索引表）
            if m and md_dir:
                full = resolve_fig(m.group("path"), md_dir)
            if full:
                cells[i].paragraphs[0].alignment = 1
                cells[i].paragraphs[0].add_run().add_picture(full, width=Inches(1.7))
                cap = cells[i].add_paragraph()
                cap.alignment = 1
                set_run(cap.add_run(cap_from_label(m.group("label"))), size=8, color=META_COLOR)
            else:
                add_rich(cells[i].paragraphs[0], c, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_fig(doc, path, caption, width=6.2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.alignment = 1
    p.add_run().add_picture(path, width=Inches(width))
    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(10)
    cap.alignment = 1
    set_run(cap.add_run(caption), size=9, color=META_COLOR)


# ── Markdown 解析 ───────────────────────────────────────────────────────────

FIG_RE = re.compile(r"^(?:📎\s*)?!?\[(?P<label>[^\]]*)\]\((?P<path>[^)]+\.png)\)\s*$")
# 行内图：行首可带一个短前缀（如 `**波形**：`），链接后面还能接说明文字
FIG_INLINE_RE = re.compile(
    r"^(?P<pre>[^📎]{0,24}?)(?P<sep>[：:]?)\s*📎\s*\[(?P<label>[^\]]*)\]\((?P<path>[^)]+\.png)\)"
    r"\s*(?P<post>.*)$")
BULLET_RE = re.compile(r"^(?P<indent>\s*)[-*]\s+(?P<text>.*)$")
NUM_RE = re.compile(r"^(?P<indent>\s*)(?P<num>\d+)\.\s+(?P<text>.*)$")
HEAD_RE = re.compile(r"^(?P<hashes>#{1,4})\s+(?P<text>.*)$")
CELL_SPLIT_RE = re.compile(r"(?<!\\)\|")
SEP_ROW_RE = re.compile(r"^\|?[\s:|-]+\|[\s:|-]*$")


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip().replace("\\|", "|") for c in CELL_SPLIT_RE.split(s)]


def is_table_start(lines, i):
    return ("|" in lines[i] and i + 1 < len(lines)
            and "|" in lines[i + 1] and SEP_ROW_RE.match(lines[i + 1].strip()))


def resolve_fig(path, md_dir):
    cand = os.path.normpath(os.path.join(md_dir, path))
    return cand if os.path.exists(cand) else None


def cap_from_label(label):
    """图注规范化：去掉「打开：」前缀与文件名后缀。

    图：门限与计数（busstate_connect_threshold.png）→ 图：门限与计数
    打开：位填充 → 位填充
    """
    s = re.sub(r"[（(][^（）()]*\.png[)）]\s*$", "", label).strip()
    return re.sub(r"^打开[：:]\s*", "", s).strip()


def strip_lead_sep(text):
    """去掉行内图后面那段说明开头的连接符（—— / — / ：）。"""
    s = text.strip()
    while s[:1] in ("—", "：", ":"):
        s = s[1:].strip()
    return s


def parse_blocks(text):
    """把 Markdown 拆成 (kind, payload) 序列。"""
    lines = text.split("\n")
    blocks, i = [], 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if not line.strip():
            i += 1
            continue

        if line.strip() in ("---", "***", "___"):
            i += 1
            continue

        m = HEAD_RE.match(line)
        if m:
            blocks.append(("head", (len(m.group("hashes")), m.group("text").strip())))
            i += 1
            continue

        if line.strip().startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1
            blocks.append(("code", code))
            continue

        if line.lstrip().startswith("> "):
            quote = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote.append(lines[i].lstrip()[1:].strip())
                i += 1
            blocks.append(("quote", [q for q in quote if q]))
            continue

        if is_table_start(lines, i):
            header = split_row(lines[i])
            i += 2
            rows = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append(("table", (header, rows)))
            continue

        fm = FIG_RE.match(line.strip())
        if fm:
            blocks.append(("fig", (fm.group("path"), cap_from_label(fm.group("label")))))
            i += 1
            continue

        im = FIG_INLINE_RE.match(line.strip())
        if im:
            pre = im.group("pre").strip("*` ").strip()
            # 只认「行首短前缀 + 图」这一种（如 `**波形**：📎 …`）；
            # 图链接夹在句子中间的长句（如「…（📎 [x.png](…)，说明）」）保持原样当段落走。
            if pre == "" or im.group("sep"):
                if pre:
                    blocks.append(("para", im.group("pre").strip()))
                blocks.append(("fig", (im.group("path"), cap_from_label(im.group("label")))))
                post = strip_lead_sep(im.group("post"))
                if post:
                    blocks.append(("para", post))
                i += 1
                continue

        bm = BULLET_RE.match(line)
        if bm:
            blocks.append(("bullet", (len(bm.group("indent")), bm.group("text"))))
            i += 1
            continue

        nm = NUM_RE.match(line)
        if nm:
            blocks.append(("num", (len(nm.group("indent")), nm.group("num"), nm.group("text"))))
            i += 1
            continue

        if line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            blocks.append(("italic", line.strip("*")))
            i += 1
            continue

        blocks.append(("para", line.strip()))
        i += 1
    return blocks


# ── 组装 ────────────────────────────────────────────────────────────────────


def build_doc(blocks, md_path, quiet=False):
    md_dir = os.path.dirname(os.path.abspath(md_path))
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    title_done = False
    missing = []
    for kind, payload in blocks:
        if kind == "head":
            level, text = payload
            if level == 1 and not title_done:
                title_done = True
                p = doc.add_paragraph()
                try:
                    p.style = doc.styles["Title"]
                except KeyError:
                    pass
                p.alignment = 1
                set_run(p.add_run(text), size=20, bold=True, color="1F3864")
                continue
            add_heading(doc, text, min(level - (1 if title_done else 0), 4))
        elif kind == "quote":
            for q in payload:
                add_para(doc, q, size=9.5, space_after=2, color=META_COLOR)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif kind == "para":
            add_para(doc, payload)
        elif kind == "bullet":
            indent, text = payload
            add_para(doc, "• " + text, indent=12 + indent * 10, space_after=2)
        elif kind == "num":
            indent, num, text = payload
            add_para(doc, f"{num}. {text}", indent=12 + indent * 10, space_after=2)
        elif kind == "italic":
            add_para(doc, payload, size=9.5, italic=True, color=META_COLOR)
        elif kind == "code":
            add_code_block(doc, payload)
        elif kind == "table":
            header, rows = payload
            add_table(doc, header, rows, md_dir=md_dir)
        elif kind == "fig":
            path, caption = payload
            full = resolve_fig(path, md_dir)
            if full:
                add_fig(doc, full, caption)
            else:
                missing.append(path)
                add_para(doc, f"[缺图：{path}]", color="C00000")
    return doc, missing


def main():
    ap = argparse.ArgumentParser(description="Markdown 报告 → Word (.docx)")
    ap.add_argument("md", help="输入的 Markdown 文件")
    ap.add_argument("-o", "--out", help="输出的 .docx（默认与 Markdown 同名同目录）")
    ap.add_argument("--dry-run", action="store_true", help="只打印解析结果，不写文件")
    args = ap.parse_args()

    text = open(args.md, encoding="utf-8").read()
    blocks = parse_blocks(text)
    doc, missing = build_doc(blocks, args.md)

    if args.dry_run:
        from collections import Counter
        print(f"{args.md}: {len(blocks)} 个块 {dict(Counter(k for k, _ in blocks))}")
        for p in missing:
            print("  缺图:", p)
        return

    out = args.out or os.path.splitext(args.md)[0] + ".docx"
    doc.save(out)
    from collections import Counter
    img = sum(1 for k, _ in blocks if k == "fig")
    tbl = sum(1 for k, _ in blocks if k == "table")
    print(f"{args.md} → {out}")
    print(f"  {len(blocks)} 块：{dict(Counter(k for k, _ in blocks))}")
    print(f"  内嵌图 {img - len(missing)}/{img}，表格 {tbl}")
    if missing:
        print("  缺图:", ", ".join(missing))


if __name__ == "__main__":
    main()
